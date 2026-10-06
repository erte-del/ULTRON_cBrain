"""read_upload: open a file the user uploaded from their computer.

text_of and jpeg_of also read files on the Mac for mac_read what=content."""

import asyncio
import base64
import io
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
from claude_agent_sdk import tool
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from storage import image_store, upload_store

MAX_CHARS = 100_000  # ~25K tokens
SCREEN_PX = 1600  # long side of a screen capture: big enough to read small text
# macOS's textutil turns these into plain text.
TEXTUTIL = {".doc", ".docx", ".rtf", ".odt", ".html", ".htm", ".webarchive", ".wordml"}
IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".heic", ".webp", ".tiff", ".bmp"}
# Pages, Numbers and Keynote files: the app exports a Word, Excel or PowerPoint copy, which is read.
IWORK = {".pages": ("Pages", "Microsoft Word", ".docx"), ".numbers": ("Numbers", "Microsoft Excel", ".xlsx"),
         ".key": ("Keynote", "Microsoft PowerPoint", ".pptx")}
EXPORT = """on run argv
  set wasRunning to application "{app}" is running
  tell application "{app}"
    set d to open (POSIX file (item 1 of argv))
    export d to (POSIX file (item 2 of argv)) as {fmt}
    close d saving no
    if not wasRunning then quit
  end tell
end run"""
READABLE = "text, code, CSV, JSON, PDF, Word, Excel, PowerPoint, Pages, Numbers, Keynote, RTF, ODT, HTML and images"


def _screen_jpeg(path: Path) -> bytes | None:
    """A screen capture (see main.screen) shrunk for Claude, or None if this isn't one."""
    if path.name != upload_store.SCREEN_NAME:
        return None
    return jpeg_of(path)


def jpeg_of(path: Path) -> bytes | None:
    """An image shrunk for Claude, or None if Pillow can't open it."""
    try:
        img = Image.open(path).convert("RGB")
    except (UnidentifiedImageError, OSError):
        return None
    img.thumbnail((SCREEN_PX, SCREEN_PX))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=85)
    return out.getvalue()


def text_of(path: Path) -> str:
    """The text in a file, cut off at MAX_CHARS. ValueError for formats it can't read."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = PdfReader(path).pages
        text = "\n\n".join(f"--- page {i} ---\n{p.extract_text() or ''}" for i, p in enumerate(pages, 1))
    elif suffix == ".pptx":
        slides = Presentation(str(path)).slides
        text = "\n\n".join(f"--- slide {i} ---\n" + "\n".join(sh.text_frame.text.replace("\v", "\n") for sh in s.shapes if sh.has_text_frame)
                           for i, s in enumerate(slides, 1))
    elif suffix in (".xlsx", ".xlsm"):
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)  # values, not formulas
        try:
            text = "\n\n".join(f"--- sheet {ws.title} ---\n" + "\n".join(
                "\t".join("" if v is None else str(v) for v in row).rstrip("\t")
                for row in ws.iter_rows(values_only=True) if any(v is not None for v in row))
                for ws in book.worksheets)
        finally:
            book.close()
    elif suffix in IWORK:
        app, fmt, ext = IWORK[suffix]
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp, path.stem + ext)
            done = subprocess.run(["osascript", "-e", EXPORT.format(app=app, fmt=fmt), str(path), str(copy)],
                                  capture_output=True, timeout=120)
            if done.returncode or not copy.exists():
                err = done.stderr.decode(errors="replace").strip()
                if "-1743" in err or "not allowed" in err.lower():
                    err = f"macOS didn't allow it. Allow Ultron to control {app} in System Settings → Privacy & Security → Automation."
                raise ValueError(f"{app} couldn't open {path.name}: {err or 'no copy came back'}")
            return text_of(copy)
    elif suffix in TEXTUTIL:
        done = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(path)], capture_output=True, timeout=60)
        if done.returncode:
            raise ValueError(f"Couldn't read {path.name}: {done.stderr.decode(errors='replace').strip()}")
        text = done.stdout.decode("utf-8", errors="replace")
    else:
        data = path.read_bytes()
        if b"\0" in data[:4096]:
            raise ValueError(f"{path.name} is a binary file; only {READABLE} can be read.")
        text = data.decode("utf-8", errors="replace")
    if len(text) > MAX_CHARS:
        # ponytail: long files are cut off; add an offset parameter if that bites
        text = text[:MAX_CHARS] + f"\n\n[cut off: only the first {MAX_CHARS} of {len(text)} characters]"
    return text


@tool(
    "read_upload",
    "Open a file the user uploaded (id like upl_003, or img_007 for an uploaded image). "
    f"{READABLE.removesuffix(' and images')} come back as text; images come back as a picture.",
    {"type": "object", "properties": {"id": {"type": "string"}}, "required": ["id"]},
)
async def read_upload(args: dict[str, Any]) -> dict[str, Any]:
    file_id = str(args["id"]).strip()
    try:
        if image_store.IMAGE_ID.match(file_id):
            rec = image_store.load(file_id)
            data = base64.b64encode(await asyncio.to_thread(image_store.preview_jpeg, rec)).decode()
            return {"content": [
                {"type": "text", "text": f"{rec.id}: {rec.title!r}"},
                {"type": "image", "data": data, "mimeType": "image/jpeg"},
            ]}
        path = upload_store.find(file_id)
        if shot := await asyncio.to_thread(_screen_jpeg, path):
            return {"content": [
                {"type": "text", "text": f"{file_id}: a capture of the user's screen, taken when they sent their message."},
                {"type": "image", "data": base64.b64encode(shot).decode(), "mimeType": "image/jpeg"},
            ]}
        text = await asyncio.to_thread(text_of, path)
    except (KeyError, ValueError, OSError, PdfReadError) as e:
        return {"content": [{"type": "text", "text": f"Couldn't open {file_id}: {e}"}], "is_error": True}
    return {"content": [{"type": "text", "text": f"{path.name}:\n\n{text}"}]}
